"""Build the six-slide SIH pitch with editable diagrams and a matching PDF.

The supplied SIH template supplies the six section layouts and official logo.
Results are read from the frozen reporting artifact, never from training scores.
Requires python-pptx, ReportLab and the Noto Sans fonts declared below.
"""
from io import BytesIO
from copy import deepcopy
from math import atan2, cos, sin
from pathlib import Path
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
INK, BLUE, MUTED = '172B3A', '285A7A', '526676'
LINE, WHITE, TEAL, AMBER = 'D7E1E7', 'FFFFFF', '27766F', '8B681F'
PALE, SOFT, NAVY = 'EAF1F5', 'F6F8FA', '162E40'
WIDTH, HEIGHT = 960, 540
REPO_URL = 'https://github.com/Siwach777/spectra-scheduler'
DEMO_URL = 'https://youtu.be/G0SuhfmfhfM?si=QQJBI2hEc8rlcnzo'
COVER_FIGURE_X, COVER_FIGURE_Y, COVER_FIGURE_SCALE = 560, 118, .88
for name, suffix in [('Regular', 'Regular'), ('Bold', 'Bold')]:
    pdfmetrics.registerFont(TTFont(name, f'/usr/share/fonts/noto/NotoSans-{suffix}.ttf'))

prs = Presentation(ROOT / 'SIH2026-IDEA-Presentation-Format.pptx')
format_reference = Presentation(ROOT / 'SIH2026-IDEA-Presentation-Format.pptx')
format_cover = format_reference.slides[0]
reference_footer = next(s for s in format_reference.slides[1].shapes
                        if s.shape_type == 1 and s.top > Pt(490))
reference_figure = next(s for s in format_cover.shapes if s.shape_type == 5)
reference_emblem = max((s for s in format_cover.shapes if s.shape_type == 13), key=lambda s: s.top)
logo = next(s.image.blob for s in prs.slides[1].shapes if s.shape_type == 13)
cover_data = json.loads((ROOT / 'assets/cover-fields.json').read_text())
cover_heading, cover_fields = cover_data['heading'], cover_data['fields']
cover_logo = (ROOT / 'assets/cover-logo.png').read_bytes()
assert len(cover_fields) == 6, 'Cover must supply all six statement/team fields'
prs.slide_width, prs.slide_height = Pt(WIDTH), Pt(HEIGHT)
while len(prs.slides) > 6:
    item = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(item.rId)
    prs.slides._sldIdLst.remove(item)
pdf = canvas.Canvas(str(ROOT / 'Spectra-Scheduler-SIH2026.pdf'), pagesize=(WIDTH, HEIGHT), invariant=1)
pdf.setTitle('Spectra Scheduler | Axiom X | SIH26055')
pdf.setAuthor('Axiom X')
slide = None
text_bounds = []


def box(x, y, w, h, fill=WHITE, border=None, rounded=False, ellipse=False):
    kind = MSO_SHAPE.OVAL if ellipse else (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
    shape = slide.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if rounded:
        shape.adjustments[0] = .1
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border:
        shape.line.color.rgb = RGBColor.from_string(border)
        shape.line.width = Pt(.8)
    else:
        shape.line.fill.background()
    pdf.setFillColor(HexColor('#' + fill))
    pdf.setStrokeColor(HexColor('#' + (border or fill)))
    pdf.setLineWidth(.8)
    if ellipse:
        pdf.ellipse(x, HEIGHT-y-h, x+w, HEIGHT-y, fill=1, stroke=bool(border))
    elif rounded:
        pdf.roundRect(x, HEIGHT-y-h, w, h, min(7, h/8), fill=1, stroke=bool(border))
    else:
        pdf.rect(x, HEIGHT-y-h, w, h, fill=1, stroke=bool(border))
    return shape


def txt(x, y, w, value, size=15, color=INK, bold=False, url=None, align='left', leading=1.3):
    lines, font = value.split('\n'), 'Bold' if bold else 'Regular'
    for content in lines:
        measured = pdfmetrics.stringWidth(content, font, size)
        if measured > w - 2:
            raise ValueError(f'Text exceeds width {w}: {content!r} ({measured:.1f} pt)')
    h = len(lines)*size*leading+4
    assert x >= 0 and x+w <= WIDTH and y >= 0 and y+h <= HEIGHT, value
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = shape.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, content in enumerate(lines):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(size*leading)
        paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
        run = paragraph.add_run()
        run.text = content
        run.font.name, run.font.size, run.font.bold = 'Noto Sans', Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        if url:
            run.hyperlink.address = url
        pdf.setFont(font, size)
        pdf.setFillColor(HexColor('#' + color))
        baseline = HEIGHT-y-size-i*size*leading
        if align == 'center':
            pdf.drawCentredString(x+w/2, baseline, content)
        elif align == 'right':
            pdf.drawRightString(x+w, baseline, content)
        else:
            pdf.drawString(x, baseline, content)
    if url:
        pdf.linkURL(url, (x, HEIGHT-y-h, x+w, HEIGHT-y))
    text_bounds.append((len(text_bounds), x, y, w, h, value))
    return shape


def line(x1, y1, x2, y2, color=LINE, width=1.3, arrow=False, dashed=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        end = OxmlElement('a:tailEnd')
        end.set('type', 'triangle')
        shape.line._get_or_add_ln().append(end)
    pdf.setStrokeColor(HexColor('#' + color))
    pdf.setLineWidth(width)
    pdf.setDash(4, 3) if dashed else pdf.setDash()
    pdf.line(x1, HEIGHT-y1, x2, HEIGHT-y2)
    pdf.setDash()
    if arrow:
        angle = atan2(y2-y1, x2-x1)
        path = pdf.beginPath()
        path.moveTo(x2, HEIGHT-y2)
        for a in (angle+.42, angle-.42):
            path.lineTo(x2-7*cos(a), HEIGHT-(y2-7*sin(a)))
        path.close()
        pdf.setFillColor(HexColor('#' + color))
        pdf.drawPath(path, fill=1, stroke=0)
    return shape


def begin(n, section, title=None, sub=None, cover=False):
    global slide
    slide = prs.slides[n-1]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    # Suppress unused template master decorations while retaining its layouts.
    slide._element.set('showMasterSp', '0')
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    box(0, 0, WIDTH, HEIGHT)
    if not cover:
        team_x, team_y, team_w, team_h = 40, 22, 108, 40
        team_badge = box(team_x, team_y, team_w, team_h, PALE, LINE, rounded=True)
        team_frame = team_badge.text_frame
        team_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        team_frame.word_wrap = False
        team_frame.margin_left = team_frame.margin_right = 0
        team_frame.margin_top = team_frame.margin_bottom = 0
        team_paragraph = team_frame.paragraphs[0]
        team_paragraph.alignment = PP_ALIGN.CENTER
        team_paragraph.space_before = team_paragraph.space_after = Pt(0)
        team_run = team_paragraph.add_run()
        team_run.text = 'Axiom X'
        team_run.font.name, team_run.font.size, team_run.font.bold = 'Noto Sans', Pt(16), True
        team_run.font.color.rgb = RGBColor.from_string(BLUE)
        pdf.setFont('Bold', 16)
        pdf.setFillColor(HexColor('#' + BLUE))
        pdf.drawCentredString(team_x+team_w/2, HEIGHT-team_y-team_h/2-5, 'Axiom X')
    txt(188, 34, 584, section, 17, BLUE, True, align='center')
    reference_logo = min((s for s in format_cover.shapes if s.shape_type == 13), key=lambda s: s.top)
    logo_x, logo_y, logo_w, logo_h = [v/12700 for v in
                                    (reference_logo.left, reference_logo.top,
                                     reference_logo.width, reference_logo.height)]
    page_logo = cover_logo if cover else logo
    slide.shapes.add_picture(BytesIO(page_logo), Pt(logo_x), Pt(logo_y), width=Pt(logo_w), height=Pt(logo_h))
    pdf.drawImage(ImageReader(BytesIO(page_logo)), logo_x, HEIGHT-logo_y-logo_h, logo_w, logo_h, mask='auto')
    line(40, 88, 920, 88, LINE, .8)
    if title:
        txt(40, 96, 880, title, 26, INK, True)
    if sub:
        txt(40, 136, 880, sub, 12, MUTED)
    if not cover:
        clone_reference_shape(reference_footer)
        footer_x, footer_y, footer_w, footer_h = [v/12700 for v in
                                                (reference_footer.left, reference_footer.top,
                                                 reference_footer.width, reference_footer.height)]
        pdf.setFillColor(HexColor('#0070C0'))
        pdf.rect(footer_x, HEIGHT-footer_y-footer_h, footer_w, footer_h, fill=1, stroke=0)
        txt(40, 514, 430, 'SIH 2026 • IDEA SUBMISSION', 8.5, WHITE, True)
        txt(844, 514, 76, f'{n}/6', 8.5, WHITE, align='right')


def end(notes):
    slide.notes_slide.notes_text_frame.text = notes
    pdf.showPage()


def node(x, y, w, title, detail, number=None, accent=BLUE, height=103):
    box(x, y, w, height, SOFT, LINE, rounded=True)
    box(x, y+16, 3, height-32, accent)
    if number:
        txt(x+15, y+12, w-30, number, 9, accent, True)
        title_y, detail_y = y+33, y+64
    else:
        title_y, detail_y = y+16, y+47
    txt(x+15, title_y, w-30, title, 17, INK, True)
    txt(x+15, detail_y, w-30, detail, 13, MUTED)


def tag(x, y, w, title, fill=PALE, color=BLUE):
    box(x, y, w, 25, fill, rounded=True)
    txt(x+8, y+5, w-16, title, 9, color, True, align='center')


def info_card(x, y, title, label, detail, height=121, title_size=16):
    """Consistent three-column content cards across idea, feasibility and impact."""
    box(x, y, 280, height, SOFT, LINE, rounded=True)
    box(x+13, y+14, 21, 2, TEAL)
    txt(x+42, y+10, 225, label, 9, TEAL, True)
    txt(x+13, y+37, 254, title, title_size, INK, True)
    txt(x+13, y+67, 254, detail, 13, MUTED, leading=1.35)


def feedback(left, right, top, y, label, size=10.5):
    """Split the return connector at the actual label width, with a 6 pt gap."""
    width = pdfmetrics.stringWidth(label, 'Regular', size)
    x = (left + right - width) / 2
    line(right, top, right, y, BLUE, 1.4)
    line(right, y, x+width+6, y, BLUE, 1.4)
    line(x-6, y, left, y, BLUE, 1.4)
    line(left, y, left, top, BLUE, 1.4, arrow=True)
    txt(x, y-size*.63, width+3, label, size, BLUE)


def clone_reference_shape(source_shape):
    """Copy a native shape, preserving the supplied geometry and appearance."""
    element = deepcopy(source_shape._element)
    element.xpath('.//p:cNvPr')[0].set('id', str(slide.shapes._next_shape_id))
    if source_shape.shape_type == 13:
        _, image_rid = slide.part.get_or_add_image_part(BytesIO(source_shape.image.blob))
        element.xpath('.//a:blip')[0].set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', image_rid)
    slide.shapes._spTree.insert_element_before(element, 'p:extLst')
    return slide.shapes[-1]


def place_cover_shape(source_shape):
    """Move and uniformly scale the complete reference figure below the header."""
    shape = clone_reference_shape(source_shape)
    shape.left = Pt(COVER_FIGURE_X + (source_shape.left-reference_figure.left)/12700*COVER_FIGURE_SCALE)
    shape.top = Pt(COVER_FIGURE_Y + (source_shape.top-reference_figure.top)/12700*COVER_FIGURE_SCALE)
    shape.width = Pt(source_shape.width/12700*COVER_FIGURE_SCALE)
    shape.height = Pt(source_shape.height/12700*COVER_FIGURE_SCALE)
    return shape


def cover_decoration():
    """Retain the supplied cover's editable freeform and mirror it in the PDF."""
    source_shape = reference_figure
    shape = place_cover_shape(source_shape)
    geometry = source_shape._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}custGeom')
    paths = geometry.find('{http://schemas.openxmlformats.org/drawingml/2006/main}pathLst')
    x, y, w, h = [v/12700 for v in (shape.left, shape.top, shape.width, shape.height)]
    pdf.saveState()
    pdf.setFillColor(HexColor('#7F7F7F'))
    pdf.setFillAlpha(.15)
    for source_path in paths:
        scale_x, scale_y = w/int(source_path.get('w')), h/int(source_path.get('h'))
        path = pdf.beginPath()
        for command in source_path:
            kind = command.tag.rsplit('}', 1)[-1]
            points = [(x+int(p.get('x'))*scale_x, HEIGHT-y-int(p.get('y'))*scale_y) for p in command]
            if kind == 'moveTo':
                path.moveTo(*points[0])
            elif kind == 'lnTo':
                path.lineTo(*points[0])
            elif kind == 'cubicBezTo':
                path.curveTo(*(v for point in points for v in point))
            elif kind == 'close':
                path.close()
            else:
                raise ValueError(f'Unsupported cover geometry: {kind}')
        pdf.drawPath(path, fill=1, stroke=0)
    pdf.restoreState()


def cover_emblem():
    """Preserve the template picture and its crop without altering the raster."""
    shape = place_cover_shape(reference_emblem)
    x, y, w, h = [v/12700 for v in (shape.left, shape.top, shape.width, shape.height)]
    left, right = shape.crop_left, shape.crop_right
    top, bottom = shape.crop_top, shape.crop_bottom
    full_w, full_h = w/(1-left-right), h/(1-top-bottom)
    pdf.saveState()
    clip = pdf.beginPath()
    clip.rect(x, HEIGHT-y-h, w, h)
    pdf.clipPath(clip, stroke=0, fill=0)
    pdf.drawImage(ImageReader(BytesIO(reference_emblem.image.blob)),
                  x-left*full_w, HEIGHT-y-full_h+top*full_h,
                  full_w, full_h, mask='auto')
    pdf.restoreState()


# Published summaries keep rebuilding independent of local training artifacts.
synthetic = json.loads((PROJECT / 'reports/timing-selected-summary.json').read_text())
external = json.loads((PROJECT / 'reports/timing-pdw-summary.json').read_text())
latency = json.loads((PROJECT / 'reports/timing-latency-summary.json').read_text())
selected = synthetic['means']['timing-trained']
pdw = external['summary_vs_sweep']
external_capture = pdw['timing-reference']['interception_ratio']['mean'] * 100
sweep_capture = pdw['sweep-50']['interception_ratio']['mean'] * 100
external_discovery = pdw['timing-reference']['discovery_fraction']['mean'] * 100
external_gain = external_capture / sweep_capture
assert len(external['plan']['recordings']) == 32
assert pdw['timing-reference']['interception_ratio']['paired_bootstrap_95_interval'][0] > 0
assert external['checkpoints']['timing-reference']['sha256'] == synthetic['frozen']['checkpoint_sha256']

begin(1, cover_heading, cover=True)
cover_decoration()
# Follow the reference's single text block: every field has the same typography.
for y, label in [(151, 'Problem Statement ID'), (201, 'Problem Statement Title'),
                 (283, 'Theme'), (327, 'PS Category'), (371, 'Team ID'), (415, 'Team Name')]:
    value = cover_fields[label]
    if label == 'Problem Statement Title':
        field = f'{label} –\n{value}'
    else:
        field = f'{label} – {value}'
    txt(40, y, 480, field, 20, INK, True)
cover_emblem()
end('Passive, receive-only software for SIH26055. Team Axiom X, ID 157808, theme Robotics and Drones and the official SIH template are retained. Requirements authority: docs/project_scope.md.')

begin(2, 'IDEA TITLE', 'Intercept the right band at the right time.')
txt(40, 141, 880, 'Spectra Scheduler learns from receiver hits and misses to decide which frequency band to observe,\nwhen to revisit it, and how long to listen—without reliable prior emitter intelligence.', 14, INK)
for x, label, title, detail in [
    (40, 'CHALLENGE', 'Fixed sweeps miss windows', 'Wide spectrum; narrow receiver.\nBrief activity between revisits.'),
    (340, 'SOLUTION', 'Learn from hits and misses', 'Forecast activity from scan history.\nAdapt band choice and dwell.'),
    (640, 'DIFFERENTIATOR', 'Coordinate frequency and time', 'Timing forecasts guide listening.\nRetune cost and coverage matter.')]:
    info_card(x, 195, title, label, detail)
for x, title, detail in [
    (40, 'Observe', 'Listening mask + detections'),
    (266, 'Predict', 'Activity and intercept time'),
    (492, 'Schedule', 'Band + listening dwell'),
    (718, 'Listen', 'Retune, then hold')]:
    node(x, 342, 202, title, detail, height=91)
for edge in (242, 468, 694):
    line(edge+1, 386, edge+23, 386, BLUE, 1.7, arrow=True)
feedback(141, 819, 434, 462, 'Learn from each revisit', 13)
end('The frozen timing forecaster uses receiver-visible masks, counts and calibrated power when available. The planner chooses band and native listening dwell from observations and public receiver settings. Synthetic simulation and external synthetic TSRD replay share this causal timing path through timing_replay.py. Replay explicitly marks calibrated power unavailable and retains raw delivered counts. Model weights remain frozen during evaluation; labels, emitter identity, known periods/phases and future truth are excluded from runtime inputs. The GUI currently displays synthetic comparisons; the replay CLI runs recordings independently.')

begin(3, 'TECHNICAL APPROACH', 'Predict activity. Choose the next scan.',
      'Selected scheduler: causal temporal forecasts with retune-aware planning.')
txt(40, 164, 880, 'Repository: github.com/Siwach777/spectra-scheduler', 11.5, BLUE, url=REPO_URL)
for x, title, detail in [
    (40, 'Receiver history', 'Previous 288 ms\nMask + observed counts'),
    (266, 'Forecast', 'Next 80 ms / all bands\nTemporal CNN + phase mix'),
    (492, 'Scan planner', 'Retune-aware planning\nBand + 1 / 10 / 50-ms dwell'),
    (718, 'Receiver', 'Simulation or PDW replay\nListen, then update history')]:
    node(x, 191, 202, title, detail, height=100)
for edge in (242, 468, 694):
    line(edge+1, 239, edge+23, 239, BLUE, 1.6, arrow=True)
feedback(141, 819, 292, 316, 'Hits and misses update the model state', 12.5)
txt(40, 348, 880, 'OUTPUTS   Capture probability • intercept time • interception ratio', 13, BLUE, True)
txt(40, 383, 880, 'CAUSAL INPUTS   No emitter identity • no supplied period / phase • no future information', 12, MUTED)
line(40, 417, 920, 417, LINE, .8)
for x, title, detail in [
    (40, 'Simulation + learning', 'Python / NumPy\nPyTorch / CUDA'),
    (340, 'Pulse replay + data', 'HDF5 / h5py\nscikit-learn'),
    (640, 'API + visualization', 'JSON / HTTP\nJavaScript / Canvas')]:
    txt(x, 435, 280, title, 14, INK, True)
    txt(x, 460, 280, detail, 12, MUTED, leading=1.2)
end('Selected checkpoint: artifacts/timing-refine-v1/seed-0/best.pt, epoch 24; SHA256 f821a7672b5b378154caa89d981893b77a62ddc0119ea74053483aa611ca744f. The band-shared TCN learns generic timing hypotheses (2–144 ms) and an aperiodic expert from 288 ms of causal receiver history, forecasting the next 80 ms. These are learned hypotheses, not supplied emitter intelligence. Receding-horizon planning accounts for retune delay, executes the first action and replans after feedback. Eight bands times 1/10/50-ms dwell gives 24 actions. Age-based coverage probes remain enabled. The synthetic CLI and browser load this selected forecaster; the replay CLI uses the corrected missing-power adapter. CUDA training, bounded memmaps, pinned staging, batched inference and configurable workers are implemented. Trajectory RL, neural MPC and other learning branches are comparison or experimental methods; the major learned contribution is temporal forecasting. External intercept-time calibration and hardware deadlines are not established.')

begin(4, 'FEASIBILITY AND VIABILITY', 'Validated in simulation and external pulse replay.',
      'Matched receiver budgets, stronger controls and a working closed-loop demo.')
for x, label, title, detail in [
    (40, 'ENGINEERING', 'Receiver simulation', 'Agile, spatial and periodic emitters.\nNoise, sensitivity and retuning.'),
    (340, 'LEARNING', 'Trained scheduler', 'Causal forecasts and scan control.\nCUDA training; held-out checks.'),
    (640, 'DATA AND SOFTWARE', 'Replay + analyst console', '32 unseen recordings evaluated.\nReplay CLI + trained-policy GUI.')]:
    info_card(x, 165, title, label, detail)
box(40, 308, 880, 107, PALE, rounded=True)
txt(58, 321, 844, 'EXTERNAL SYNTHETIC TSRD VALIDATION • 32 RECORDINGS', 10, TEAL, True)
txt(58, 344, 844, f'{external_gain:.2f}× mean capture versus fixed sweep', 22, BLUE, True)
txt(58, 380, 844, f'Timing {external_capture:.2f}% • fixed sweep {sweep_capture:.2f}% • emitter discovery {external_discovery:.2f}%', 13, INK)
for x, title, detail in [
    (40, '300-world simulation', 'Agile / spatial / periodic capture:\n34.37% / 51.56% / 67.66%'),
    (340, 'Capture / discovery', 'Rate-probe: 51.66% capture;\n84.39% emitter discovery.'),
    (640, 'Validation boundary', 'External synthetic data.\nNo RF hardware validation.')]:
    txt(x, 444, 280, title, 14, INK, True)
    txt(x, 463, 280, detail, 12, MUTED, leading=1.2)
end('Frozen synthetic benchmark: 300 fresh worlds, 100 per emitter family, identical physical-time receiver budgets. Timing mean capture is 34.37% agile, 51.56% spatial and 67.66% periodic versus fixed sweep 11.10%, 11.17%, 10.87%. Strong phase planning achieves 27.57%, 45.85%, 64.15%; periodic superiority to phase planning is not established. Spatial discovery is 86.0% versus sweep 97.5%. External synthetic TSRD validation uses 32 previously unseen recordings, 10 seconds per recording, 8 bands, 2-ms retuning and detection probability 0.9. Corrected frozen timing capture is 43.4882% versus sweep 10.4465%, a ratio of mean recording capture of 4.16294; paired gain is 33.04 percentage points, 95% interval [28.19, 38.12]. Discovery is 93.35% versus rate-probe 84.39%, while rate-probe capture remains higher at 51.66%. The legacy adapter gave 9.85% capture. The improvement fixes missing-power gating and clipped counts; it is not a new neural architecture. A CUDA fine-tune on 64 train recordings with 16 disjoint train-split development recordings reached 42.72% on the new validation set and did not improve the corrected frozen model. The test split was untouched. TSRD is external synthetic data, not real RF. Sources: reports/timing-selected-summary.json and reports/timing-pdw-summary.json.')

begin(5, 'IMPACT AND BENEFITS', 'Adaptive surveillance. Measurable interception.',
      'Implemented scheduling, prediction and evaluation for Electronic Support receivers.')
for x, label, title, detail in [
    (40, 'SURVEILLANCE', 'Capture brief opportunities', 'Forecast activity before revisits.\nAdapt to agile and periodic emitters.'),
    (340, 'RECEIVER RESOURCES', 'Use receiver time well', 'Choose band and dwell together.\nAccount for retuning and coverage.'),
    (640, 'DEVELOPMENT', 'Evaluate faster', 'Reuse RF worlds and pulse replay.\nCompare policies on matched data.')]:
    info_card(x, 173, title, label, detail, height=128)
txt(40, 334, 880, 'CORE CAPABILITIES IMPLEMENTED', 10, TEAL, True)
for x, title, detail in [
    (40, 'Learned scan scheduling', 'Band + dwell control from causal receiver history.\nIntercept-time and interception-ratio predictions.'),
    (500, 'Simulation, replay and evaluation', 'Agile, spatial and periodic emitter scenarios.\nRequired figures of merit + physical-time replay.')]:
    txt(x, 366, 420, title, 18, INK, True)
    txt(x, 400, 420, detail, 13, MUTED)
line(40, 456, 920, 456, LINE, .8)
txt(40, 475, 880, 'Warmed serial scheduler p99: 0.464–0.532 ms on CUDA; receiver I/O and cold startup excluded.', 13, MUTED)
end('The browser provides synchronized receiver views, historical detections, causal future opportunities, selected band/dwell and live capture/discovery metrics with JSON export. A selected periodic seed-42 video example captures 62/86 versus fixed sweep 13/86, or 4.77 times, with both discovering all emitters. This is one illustrative simulation, separate from aggregate benchmarks. Optional compiled planning plus captured CUDA inference recorded warmed serial full scheduler p99 0.464–0.532 ms over 5986 decisions on the RTX 5070 Laptop GPU, below a declared 1-ms software budget. The measured path includes feedback ingestion, history encoding, neural inference, planning and forecast creation. It excludes checkpoint loading, warmup, simulator truth generation, receiver I/O and GUI serialization; it is not a hardware deadline guarantee. CI verifies correctness lint, CPU tests, HTTP/frontend checks and a deterministic CLI smoke benchmark. Coverage acquisition remains experimental: fresh spatial discovery increased 81.0% to 83.5%, but the paired 95% interval [-2.5, 7.5] percentage points crosses zero. Sources: reports/timing-latency-summary.json, reports/timing-acquisition-summary.json, web/README.md and docs/timing-model-findings.md.')

begin(6, 'RESEARCH AND REFERENCES', 'Research, source code and demonstration.',
      'Primary research sources and project links.')
refs = [
    ('01','Radar data','Gunn et al. — The Turing Synthetic Radar Dataset','Pulse descriptor words for association and replay.','arXiv:2602.03856','https://arxiv.org/abs/2602.03856'),
    ('02','Dataset contract','The Alan Turing Institute — TSRD dataset card','Receiver modes, feature units and data splits.','Official dataset card','https://huggingface.co/datasets/alan-turing-institute/turing-synthetic-radar-dataset'),
    ('03','Temporal prediction','Bai, Kolter & Koltun — Sequence Modeling','Causal temporal convolution for activity forecasting.','arXiv:1803.01271','https://arxiv.org/abs/1803.01271'),
    ('04','Policy learning','Ahmadian et al. — Back to Basics / RLOO','Trajectory-level learning for scan decisions.','arXiv:2402.14740','https://arxiv.org/abs/2402.14740'),
]
for i,(number,role,title,detail,label,url) in enumerate(refs):
    y=167+i*60
    txt(40,y,42,number,23,BLUE,True)
    txt(95,y+1,239,role,14,INK,True)
    txt(355,y,565,title,14,INK,True,url=url)
    txt(355,y+27,565,detail,12,MUTED)
    txt(95,y+27,240,label,11,BLUE,url=url)
    if i<3:
        line(95,y+53,920,y+53,LINE,.6)
for x, label, display, url in [
    (40, 'SOURCE CODE • GITHUB', 'github.com/Siwach777/spectra-scheduler', REPO_URL),
    (500, 'DEMO VIDEO • YOUTUBE', 'youtu.be/G0SuhfmfhfM', DEMO_URL)]:
    box(x, 417, 420, 55, SOFT, LINE, rounded=True)
    txt(x+14, 426, 392, label, 9.5, TEAL, True)
    txt(x+14, 447, 392, display, 12, BLUE, url=url)
txt(40, 480, 880, 'METRICS   Detection • false alarms • sensitivity • intercept rate • reward / cost • accuracy • timing error', 10.5, MUTED)
end('Primary related work: TSRD paper arXiv:2602.03856 and the Alan Turing Institute dataset card; TCN sequence modeling arXiv:1803.01271; RLOO/REINFORCE learning from human feedback arXiv:2402.14740. RLOO informs an implemented trajectory-policy experiment; no LLM runs in the receiver scheduler. Project evidence: reports/timing-selected-summary.json, reports/timing-pdw-summary.json and reports/timing-latency-summary.json. docs/project_scope.md defines the requirements. Main implemented system: temporal forecaster, band/dwell planner, passive receiver simulation, PDW replay, CLI and browser demo. Results are synthetic development and external synthetic validation evidence; real RF, hardware calibration and final test reporting remain open. Repository and demonstration links are supplied project links.')

prs.core_properties.title = 'Spectra Scheduler — Smart Scan Strategy for Electronic Warfare'
prs.core_properties.author = 'Axiom X'
prs.core_properties.subject = 'SIH26055 | ML-based Electronic Support receiver scheduler'
prs.core_properties.keywords = 'Axiom X, SIH26055, smart scan, electronic support, timing forecast, receiver scheduling'
prs.save(ROOT / 'Spectra-Scheduler-SIH2026.pptx')
pdf.save()
assert len(prs.slides) == 6
assert [sum(shape.shape_type == 13 for shape in page.shapes) for page in prs.slides] == [2,1,1,1,1,1]
print('Built six slides: template-style cover, compact feedback loops, feasibility, benefits and technology stack.')
